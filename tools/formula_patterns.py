"""Collapse workbook formulas into distinct relative (R1C1-style) patterns per sheet row.

Usage: python tools/formula_patterns.py WORKBOOK.xlsx OUT.csv
"""
import csv, re, sys
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

SHEETS = ['Consumption', 'Equipment', 'IntellPropProd', 'NonresStructures', 'Residential',
          'Inventories', 'ExportsImportsGoods', 'ExportsImportsServices', 'FederalGovt', 'StateLocal']
REF = re.compile(r"(?<![A-Za-z0-9_$])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(A-Za-z_])")


def relative(formula, row, col):
    def sub(m):
        cabs, c, rabs, r = m.groups()
        cc = column_index_from_string(c); rr = int(r)
        cs = f'C{cc}' if cabs else f'C[{cc - col}]'
        rs = f'R{rr}' if rabs else f'R[{rr - row}]'
        return rs + cs
    return REF.sub(sub, formula)


def main(path, out):
    wf = openpyxl.load_workbook(path, read_only=True)
    rows_out = []
    for s in SHEETS:
        for r, row in enumerate(wf[s].iter_rows(values_only=True), 1):
            label = ' | '.join(str(x)[:50] for x in row[:5] if isinstance(x, str) and not x.startswith('='))
            pats = {}
            for c, x in enumerate(row, 1):
                if isinstance(x, str) and x.startswith('='):
                    p = relative(x, r, c)
                    pats.setdefault(p, []).append(get_column_letter(c))
            for p, cols in pats.items():
                rows_out.append([s, r, label, len(cols), f'{cols[0]}:{cols[-1]}', p])
    with open(out, 'w', newline='') as f:
        w = csv.writer(f); w.writerow(['sheet', 'row', 'label', 'ncells', 'cols', 'pattern']); w.writerows(rows_out)
    print(f'{len(rows_out)} distinct row patterns')


if __name__ == '__main__':
    main(*sys.argv[1:3])
